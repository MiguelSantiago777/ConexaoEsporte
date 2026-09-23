"""
Testes da importação em massa (planilha .xlsx) — cobre o contrato comum às
5 entidades importáveis via o caso mais completo (Beneficiários: FK por
nome, validação de menor de idade) e casos específicos de Usuários (senha
gerada + e-mail) e Produtos (caso mais simples).
"""
import io

import openpyxl
import pytest

from app.application.importacao.geocodificacao import Coordenada
from tests.conftest import login


def _planilha(cabecalho: list[str], linhas: list[list]) -> tuple[str, io.BytesIO, str]:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Dados"
    ws.append(cabecalho)
    for linha in linhas:
        ws.append(linha)
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return (
        "planilha.xlsx", buffer,
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


def _upload(client, token, url, arquivo, confirmar=False):
    return client.post(
        url,
        params={"confirmar": str(confirmar).lower()},
        files={"arquivo": arquivo},
        headers={"Authorization": f"Bearer {token}"},
    )


# ---------------------------------------------------------------------
# Beneficiários
# ---------------------------------------------------------------------


def test_modelo_beneficiarios_inclui_coluna_polo_para_master(client, seed_basico):
    token = login(client, "master@test.com")
    resp = client.get(
        "/api/v1/beneficiarios/importar/modelo", headers={"Authorization": f"Bearer {token}"}
    )
    assert resp.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(resp.content))
    cabecalho = [c.value for c in next(wb["Dados"].iter_rows(min_row=1, max_row=1))]
    assert "Polo*" in cabecalho


def test_modelo_beneficiarios_omite_coluna_polo_para_gestor(client, seed_basico):
    token = login(client, "gestor.a@test.com")
    resp = client.get(
        "/api/v1/beneficiarios/importar/modelo", headers={"Authorization": f"Bearer {token}"}
    )
    assert resp.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(resp.content))
    cabecalho = [c.value for c in next(wb["Dados"].iter_rows(min_row=1, max_row=1))]
    assert "Polo*" not in cabecalho


def test_previa_valida_sem_gravar_nada(client, seed_basico):
    token = login(client, "master@test.com")
    polo_a = seed_basico["polo_a"].nome
    arquivo = _planilha(
        ["Nome completo*", "Data de nascimento (DD/MM/AAAA)*", "Documento (CPF)*", "Polo*"],
        [
            ["Maria Teste", "10/05/2000", "111.111.111-11", polo_a],
            ["Sem Documento", "10/05/2000", "", polo_a],  # linha inválida: documento vazio
        ],
    )
    resp = _upload(client, token, "/api/v1/beneficiarios/importar", arquivo, confirmar=False)
    assert resp.status_code == 200, resp.text
    corpo = resp.json()
    assert corpo["confirmado"] is False
    assert corpo["total"] == 2
    assert corpo["sucesso"] == 1
    assert corpo["falha"] == 1
    assert corpo["linhas"][1]["erro"]

    # nada foi gravado — a lista continua vazia
    resp_lista = client.get("/api/v1/beneficiarios", headers={"Authorization": f"Bearer {token}"})
    assert resp_lista.json() == []


def test_confirmar_grava_linhas_validas_e_pula_invalidas(client, seed_basico):
    token = login(client, "master@test.com")
    polo_a = seed_basico["polo_a"].nome
    arquivo = _planilha(
        ["Nome completo*", "Data de nascimento (DD/MM/AAAA)*", "Documento (CPF)*", "Polo*"],
        [
            ["Maria Teste", "10/05/2000", "111.111.111-11", polo_a],
            ["Sem Documento", "10/05/2000", "", polo_a],
        ],
    )
    resp = _upload(client, token, "/api/v1/beneficiarios/importar", arquivo, confirmar=True)
    assert resp.status_code == 200, resp.text
    corpo = resp.json()
    assert corpo["confirmado"] is True
    assert corpo["sucesso"] == 1
    assert corpo["falha"] == 1

    resp_lista = client.get("/api/v1/beneficiarios", headers={"Authorization": f"Bearer {token}"})
    beneficiarios = resp_lista.json()
    assert len(beneficiarios) == 1
    assert beneficiarios[0]["nome_completo"] == "Maria Teste"


def test_gestor_polo_importa_direto_no_proprio_polo(client, seed_basico):
    token = login(client, "gestor.a@test.com")
    arquivo = _planilha(
        ["Nome completo*", "Data de nascimento (DD/MM/AAAA)*", "Documento (CPF)*"],
        [["João Teste", "10/05/2000", "222.222.222-22"]],
    )
    resp = _upload(client, token, "/api/v1/beneficiarios/importar", arquivo, confirmar=True)
    assert resp.status_code == 200, resp.text
    assert resp.json()["sucesso"] == 1

    resp_lista = client.get("/api/v1/beneficiarios", headers={"Authorization": f"Bearer {token}"})
    beneficiarios = resp_lista.json()
    assert len(beneficiarios) == 1
    assert beneficiarios[0]["polo_id"] == str(seed_basico["polo_a"].id)


def test_polo_nao_encontrado_da_erro_por_linha_claro(client, seed_basico):
    token = login(client, "master@test.com")
    arquivo = _planilha(
        ["Nome completo*", "Data de nascimento (DD/MM/AAAA)*", "Documento (CPF)*", "Polo*"],
        [["Maria Teste", "10/05/2000", "111.111.111-11", "Polo Que Não Existe"]],
    )
    resp = _upload(client, token, "/api/v1/beneficiarios/importar", arquivo, confirmar=False)
    assert resp.status_code == 200
    corpo = resp.json()
    assert corpo["falha"] == 1
    assert "Polo Que Não Existe" in corpo["linhas"][0]["erro"]


def test_menor_de_idade_sem_responsavel_da_erro(client, seed_basico):
    token = login(client, "master@test.com")
    polo_a = seed_basico["polo_a"].nome
    arquivo = _planilha(
        ["Nome completo*", "Data de nascimento (DD/MM/AAAA)*", "Documento (CPF)*", "Polo*"],
        [["Criança Teste", "10/05/2015", "333.333.333-33", polo_a]],
    )
    resp = _upload(client, token, "/api/v1/beneficiarios/importar", arquivo, confirmar=False)
    assert resp.status_code == 200
    corpo = resp.json()
    assert corpo["falha"] == 1
    assert "responsável" in corpo["linhas"][0]["erro"].lower()


# ---------------------------------------------------------------------
# Usuários/professores — senha gerada + e-mail de credenciais
# ---------------------------------------------------------------------


def test_importar_usuarios_gera_senha_e_envia_email_so_na_confirmacao(client, seed_basico, _sem_envio_real_de_email):
    token = login(client, "master@test.com")
    polo_a = seed_basico["polo_a"].nome
    arquivo = _planilha(
        ["Nome*", "E-mail*", "Perfil*", "Polo"],
        [["Professor Teste", "professor.teste@email.com", "PROFESSOR", polo_a]],
    )

    resp_previa = _upload(client, token, "/api/v1/usuarios/importar", arquivo, confirmar=False)
    assert resp_previa.status_code == 200, resp_previa.text
    assert resp_previa.json()["sucesso"] == 1
    assert _sem_envio_real_de_email == []  # prévia nunca envia email

    arquivo2 = _planilha(
        ["Nome*", "E-mail*", "Perfil*", "Polo"],
        [["Professor Teste", "professor.teste@email.com", "PROFESSOR", polo_a]],
    )
    resp_confirma = _upload(client, token, "/api/v1/usuarios/importar", arquivo2, confirmar=True)
    assert resp_confirma.status_code == 200, resp_confirma.text
    assert resp_confirma.json()["sucesso"] == 1
    assert len(_sem_envio_real_de_email) == 1
    assert _sem_envio_real_de_email[0]["destinatario"] == "professor.teste@email.com"


def test_gestor_polo_so_pode_importar_professor(client, seed_basico):
    token = login(client, "gestor.a@test.com")
    arquivo = _planilha(["Nome*", "E-mail*", "Perfil*"], [["Outro Master", "outromaster@email.com", "MASTER"]])
    resp = _upload(client, token, "/api/v1/usuarios/importar", arquivo, confirmar=False)
    assert resp.status_code == 200
    corpo = resp.json()
    assert corpo["falha"] == 1
    assert corpo["sucesso"] == 0


# ---------------------------------------------------------------------
# Polos — geocodificação automática do endereço
# ---------------------------------------------------------------------


def test_importar_polos_geocodifica_endereco_quando_falta_coordenada(client, seed_basico, monkeypatch):
    from app.application.polo import importacao_service as polo_importacao_service

    chamadas = []

    def _geocodificar_falso(endereco):
        chamadas.append(endereco)
        return Coordenada(-23.5, -46.6, aproximado=False)

    monkeypatch.setattr(polo_importacao_service, "geocodificar", _geocodificar_falso)

    token = login(client, "master@test.com")
    arquivo = _planilha(["Nome*", "Endereço"], [["Polo Geocodificado", "Rua Teste, 123"]])
    resp = _upload(client, token, "/api/v1/polos/importar", arquivo, confirmar=True)
    assert resp.status_code == 200, resp.text
    assert resp.json()["sucesso"] == 1
    assert chamadas == ["Rua Teste, 123"]

    resp_lista = client.get("/api/v1/polos", headers={"Authorization": f"Bearer {token}"})
    polo = next(p for p in resp_lista.json() if p["nome"] == "Polo Geocodificado")
    assert polo["latitude"] == -23.5
    assert polo["longitude"] == -46.6


def test_importar_polos_avisa_localizacao_aproximada_e_endereco_nao_encontrado(client, seed_basico, monkeypatch):
    """Antes, um endereço não encontrado deixava o polo fora do mapa em
    silêncio. Agora a linha continua OK (o polo é criado), mas com aviso."""
    from app.application.polo import importacao_service as polo_importacao_service

    respostas = {
        "Rua Achada, 1 - Bairro: Centro/RJ": Coordenada(-22.9, -43.2, aproximado=False),
        "Rua Só Bairro - Bairro: Freguesia/RJ": Coordenada(-22.94, -43.34, aproximado=True),
        "Lugar Nenhum": None,
    }
    monkeypatch.setattr(polo_importacao_service, "geocodificar", lambda endereco: respostas[endereco])

    token = login(client, "master@test.com")
    arquivo = _planilha(
        ["Nome*", "Endereço"],
        [["Polo Exato", "Rua Achada, 1 - Bairro: Centro/RJ"],
         ["Polo Aproximado", "Rua Só Bairro - Bairro: Freguesia/RJ"],
         ["Polo Sem Mapa", "Lugar Nenhum"]],
    )
    resp = _upload(client, token, "/api/v1/polos/importar", arquivo, confirmar=False)
    assert resp.status_code == 200, resp.text
    linhas = {l["resumo"]: l for l in resp.json()["linhas"]}
    assert all(l["status"] == "ok" for l in linhas.values())
    assert linhas["Polo Exato"]["aviso"] is None
    assert "aproximada" in linhas["Polo Aproximado"]["aviso"].lower()
    assert "não vai aparecer no mapa" in linhas["Polo Sem Mapa"]["aviso"]


def test_consultas_candidatas_limpam_enderecos_da_planilha_real():
    """Formatos reais que o Nominatim não achava (11 de 12 de uma planilha
    de núcleos): "Bairro:", "esquina com", "/RJ", "s/n", "Cobertura"."""
    from app.application.importacao.geocodificacao import consultas_candidatas

    def consultas(endereco):
        return [c for c, _ in consultas_candidatas(endereco)]

    assert consultas("Rua Geovani de Castro, 90  - Bairro: Freguesia/RJ")[0] == (
        "Rua Geovani de Castro, 90, Freguesia, RJ, Brasil"
    )
    assert consultas("Avenida Gerémario Dantas, 436, Cobertura - Bairro: Tanque/RJ")[0] == (
        "Avenida Gerémario Dantas, 436, Tanque, RJ, Brasil"
    )
    assert consultas(
        "Rua Professora Evangelina Guedes - esquina com R: Santa Catarina - Bairro: Cidade Luz - Campos dos Goytacazes/RJ"
    )[0] == "Rua Professora Evangelina Guedes, Cidade Luz, Campos dos Goytacazes, RJ, Brasil"
    assert consultas("Rua Ibitioca, s/n, bairro: Parque Lebret - Campos dos Goytacazes/RJ")[0] == (
        "Rua Ibitioca, Parque Lebret, Campos dos Goytacazes, RJ, Brasil"
    )
    # A última tentativa é sempre só o bairro/cidade — marcada como aproximada.
    ultima, aproximado = consultas_candidatas("Rua Geovani de Castro, 90 - Bairro: Freguesia/RJ")[-1]
    assert ultima == "Freguesia, RJ"
    assert aproximado is True


def test_localizar_polo_sem_coordenada_pelo_endereco(client, seed_basico, monkeypatch):
    from app.application.polo import service as polo_service

    monkeypatch.setattr(polo_service, "geocodificar", lambda endereco: Coordenada(-22.8, -43.4, aproximado=True))
    token = login(client, "master@test.com")
    headers = {"Authorization": f"Bearer {token}"}
    criado = client.post("/api/v1/polos", json={"nome": "Polo Sem Pino", "endereco": "Rua X - Bairro: Anchieta/RJ"}, headers=headers)
    assert criado.status_code == 201, criado.text

    resp = client.post(f"/api/v1/polos/{criado.json()['id']}/localizar", headers=headers)
    assert resp.status_code == 200, resp.text
    assert resp.json()["situacao"] == "aproximado"
    assert resp.json()["polo"]["latitude"] == -22.8

    monkeypatch.setattr(polo_service, "geocodificar", lambda endereco: None)
    resp_nada = client.post(f"/api/v1/polos/{criado.json()['id']}/localizar", headers=headers)
    assert resp_nada.json()["situacao"] == "nao_encontrado"
    assert resp_nada.json()["polo"]["latitude"] == -22.8  # não apaga o que já tinha

    token_gestor = login(client, "gestor.a@test.com")
    resp_gestor = client.post(
        f"/api/v1/polos/{criado.json()['id']}/localizar", headers={"Authorization": f"Bearer {token_gestor}"}
    )
    assert resp_gestor.status_code == 403


def test_importar_polos_nao_geocodifica_quando_coordenada_ja_vem_preenchida(client, seed_basico, monkeypatch):
    from app.application.polo import importacao_service as polo_importacao_service

    def _geocodificar_falso(endereco):
        raise AssertionError("não deveria geocodificar quando lat/long já vêm preenchidos")

    monkeypatch.setattr(polo_importacao_service, "geocodificar", _geocodificar_falso)

    token = login(client, "master@test.com")
    arquivo = _planilha(
        ["Nome*", "Endereço", "Latitude", "Longitude"],
        [["Polo Com Coordenada", "Rua Teste, 123", "-10.1", "-20.2"]],
    )
    resp = _upload(client, token, "/api/v1/polos/importar", arquivo, confirmar=True)
    assert resp.status_code == 200, resp.text
    assert resp.json()["sucesso"] == 1


# ---------------------------------------------------------------------
# Produtos — caso mais simples (sem FK, sem RBAC de polo)
# ---------------------------------------------------------------------


def test_importar_produtos(client, seed_basico):
    token = login(client, "master@test.com")
    arquivo = _planilha(
        ["Nome*", "Unidade de medida*", "Descrição"],
        [["Bola de futebol", "unidade", ""], ["Colete", "unidade", "Colete de treino"]],
    )
    resp = _upload(client, token, "/api/v1/produtos/importar", arquivo, confirmar=True)
    assert resp.status_code == 200, resp.text
    corpo = resp.json()
    assert corpo["sucesso"] == 2
    assert corpo["falha"] == 0


def test_importar_produtos_com_quantidade_e_ncm(client, seed_basico):
    """Quantidade vira a Entrada inicial no estoque; NCM aceita pontuação e
    também número (Excel sem formatação de texto perde o zero à esquerda)."""
    token = login(client, "master@test.com")
    arquivo = _planilha(
        ["Nome*", "Unidade de medida*", "Quantidade", "NCM", "Descrição"],
        [
            ["Bola de futebol", "unidade", 25, "9506.62.00", ""],
            ["Rede", "unidade", "", 1012100, ""],
            ["Cone", "unidade", 5, "123", ""],  # NCM inválido
        ],
    )
    resp = _upload(client, token, "/api/v1/produtos/importar", arquivo, confirmar=True)
    assert resp.status_code == 200, resp.text
    corpo = resp.json()
    assert corpo["sucesso"] == 2
    assert corpo["falha"] == 1

    produtos = client.get("/api/v1/produtos", headers={"Authorization": f"Bearer {token}"}).json()
    por_nome = {p["nome"]: p for p in produtos}
    assert por_nome["Bola de futebol"]["saldo_atual"] == 25
    assert por_nome["Bola de futebol"]["ncm"] == "95066200"
    assert por_nome["Rede"]["saldo_atual"] == 0
    assert por_nome["Rede"]["ncm"] == "01012100"
    assert "Cone" not in por_nome


def test_modelo_de_importacao_de_produtos_tem_quantidade_e_ncm(client, seed_basico):
    token = login(client, "master@test.com")
    resp = client.get("/api/v1/produtos/importar/modelo", headers={"Authorization": f"Bearer {token}"})
    assert resp.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(resp.content))
    cabecalho = [c.value for c in wb["Dados"][1]]
    assert cabecalho == ["Nome*", "Unidade de medida*", "Descrição", "Quantidade", "NCM"]


def test_gestor_polo_nao_pode_importar_produtos(client, seed_basico):
    token = login(client, "gestor.a@test.com")
    arquivo = _planilha(["Nome*", "Unidade de medida*", "Descrição"], [["Bola", "unidade", ""]])
    resp = _upload(client, token, "/api/v1/produtos/importar", arquivo, confirmar=True)
    assert resp.status_code == 403
